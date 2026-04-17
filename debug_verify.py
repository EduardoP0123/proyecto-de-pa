"""
UI Form
- Formateo a 3 decimales idéntico al requerimiento de la imagen.
- Fórmula matemática KW restaurada a =MAX() * factor_demanda * multiplo
"""
import os
import sys
import threading
import tkinter as tk
from pathlib import Path
import math
import re
from datetime import datetime, timedelta
from tkinter import ttk, filedialog, messagebox
from tkcalendar import DateEntry
from collections import defaultdict
import pandas as pd

try:
    from PIL import Image, ImageTk  
except Exception:
    Image = ImageTk = None

project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))
try:
    from src.ui_components import run_ui
except Exception:
    run_ui = None
from src.csv_processor import CSVProcessor


MONTH_ABBR_ES = {
    1: "ene", 2: "feb", 3: "mar", 4: "abr", 5: "may", 6: "jun",
    7: "jul", 8: "ago", 9: "sep", 10: "oct", 11: "nov", 12: "dic"
}

def format_es_date(d: datetime.date) -> str:
    return f"{d.day}-{MONTH_ABBR_ES.get(d.month, '')}-{d.strftime('%y')}"

class CSVUploaderApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Lecturas KV2C - v2.0")
        try:
            self.root.tk.call("tk", "scaling", 1.25)
        except Exception:
            pass

        self._init_style()

        self.root.geometry("980x680")
        self.root.minsize(900, 600)
        self.root.rowconfigure(0, weight=1)
        self.root.columnconfigure(0, weight=1)

        self.seg_logo_img = None

        workspace_path = Path.home() / "Downloads" / "BILLREAD_WORKSPACE"
        workspace_path.mkdir(parents=True, exist_ok=True)
        self.csv_processor = run_ui(workspace_path) if run_ui else CSVProcessor(workspace_path)

        self.create_widgets()
        self._build_statusbar()

        self.company_multipliers = {}
        self.last_report = None

    def _init_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass
        self.COLOR_BG = "#0D1B2A"
        self.COLOR_ACCENT = "#1B6F9B"
        self.COLOR_BTN = "#1F7A8C"

        style.configure("Header.TFrame", background=self.COLOR_BG)
        style.configure("Header.TLabel", background=self.COLOR_BG, foreground="white", font=("Segoe UI", 16, "bold"))
        style.configure("TLabel", font=("Segoe UI", 10))
        style.configure("TButton", font=("Segoe UI", 10))
        style.configure("Accent.TButton", font=("Segoe UI", 10, "bold"), foreground="white", background=self.COLOR_BTN)
        style.map("Accent.TButton", background=[("active", "#2390A6"), ("disabled", "#9fb6bf")])
        style.configure("Card.TFrame", padding=12)
        style.configure("Section.TLabelframe", padding=12)
        style.configure("Section.TLabelframe.Label", font=("Segoe UI", 11, "bold"))

    def _load_seg_logo(self, max_h=56, max_w=220):
        try:
            here = Path(__file__).resolve().parent
            roots = [
                here / "images",                    
                here.parent / "assets",             
                Path.cwd() / "ui" / "images",       
                Path.cwd() / "assets",
            ]
            candidates = [
                "seg.png", "SEG.png", "seg_logo.png",
                "seg.jpg", "SEG.jpg", "seg_logo.jpg",
                "seg.gif", "SEG.gif"
            ]
            for root in roots:
                for name in candidates:
                    p = root / name
                    if not p.exists():
                        continue
                    if Image and ImageTk:
                        img = Image.open(p).convert("RGBA")
                        r = min(max_h / img.height, max_w / img.width, 1.0)
                        new_size = (max(1, int(img.width * r)), max(1, int(img.height * r)))
                        img = img.resize(new_size, Image.LANCZOS)
                        return ImageTk.PhotoImage(img)
                    pic = tk.PhotoImage(file=str(p))
                    h, w = pic.height(), pic.width()
                    factor = max(1, math.ceil(max(h / max_h, w / max_w)))
                    if factor > 1:
                        pic = pic.subsample(factor, factor)
                    return pic
        except Exception as e:
            try:
                self.append_info(f"[LOGO] No se pudo cargar: {e}")
            except Exception:
                pass
        return None

    def create_widgets(self):
        root_frame = ttk.Frame(self.root, padding=0)
        root_frame.grid(row=0, column=0, sticky="nsew")
        root_frame.rowconfigure(1, weight=1)
        root_frame.columnconfigure(0, weight=1)

        header = ttk.Frame(root_frame, style="Header.TFrame", padding=(16, 12))
        header.grid(row=0, column=0, sticky="ew")
        header.columnconfigure(0, weight=1)

        ttk.Label(header, text="Lecturas KV2C / KV2A analyzer", style="Header.TLabel").grid(row=0, column=0, sticky="w")
        self.seg_logo_img = self._load_seg_logo()
        if self.seg_logo_img:
            ttk.Label(header, image=self.seg_logo_img, style="Header.TLabel").grid(row=0, column=1, sticky="e")
        else:
            ttk.Label(header, text="SEG", style="Header.TLabel").grid(row=0, column=1, sticky="e")

        body = ttk.Frame(root_frame, padding=12)
        body.grid(row=1, column=0, sticky="nsew")
        body.columnconfigure(0, weight=0)
        body.columnconfigure(1, weight=1)
        body.rowconfigure(0, weight=1)

        opts = ttk.Labelframe(body, text="Opciones", style="Section.TLabelframe")
        opts.grid(row=0, column=0, sticky="ns", padx=(0, 12))
        for i in range(6):
            opts.columnconfigure(i, weight=0)

        ttk.Label(opts, text="Carpeta").grid(row=0, column=0, sticky="w")
        self.folder_path = ttk.Entry(opts, width=40)
        self.folder_path.grid(row=0, column=1, columnspan=4, sticky="ew", padx=8)
        ttk.Button(opts, text="Examinar...", command=self.browse_folder).grid(row=0, column=5, sticky="e")

        ttk.Label(opts, text="Resolución").grid(row=1, column=0, sticky="w", pady=(8, 0))
        self.resolution = ttk.Combobox(opts, values=["15min", "1h"], state="readonly", width=8)
        self.resolution.set("15min")
        self.resolution.grid(row=1, column=1, sticky="w", padx=8, pady=(8, 0))

        ttk.Label(opts, text="Inicio").grid(row=2, column=0, sticky="w", pady=(8, 0))
        self.start_date = DateEntry(opts, date_pattern="dd/mm/y", width=10)
        self.start_date.grid(row=2, column=1, sticky="w", padx=(8, 4), pady=(8, 0))
        self.start_hour = ttk.Spinbox(opts, from_=0, to=23, width=3)
        self.start_hour.set("00")
        self.start_hour.grid(row=2, column=2, sticky="w", pady=(8, 0))
        self.start_min = ttk.Spinbox(opts, from_=0, to=59, width=3)
        self.start_min.set("00")
        self.start_min.grid(row=2, column=3, sticky="w", padx=(4, 0), pady=(8, 0))

        ttk.Label(opts, text="Fin").grid(row=3, column=0, sticky="w")
        self.end_date = DateEntry(opts, date_pattern="dd/mm/y", width=10)
        self.end_date.grid(row=3, column=1, sticky="w", padx=(8, 4))
        self.end_hour = ttk.Spinbox(opts, from_=0, to=23, width=3)
        self.end_hour.set("23")
        self.end_hour.grid(row=3, column=2, sticky="w")
        self.end_min = ttk.Spinbox(opts, from_=0, to=59, width=3)
        self.end_min.set("59")
        self.end_min.grid(row=3, column=3, sticky="w", padx=(4, 0))

        ttk.Label(opts, text="Empresa").grid(row=4, column=0, sticky="w", pady=(8, 0))
        self.company_cb = ttk.Combobox(opts, values=[], state="disabled", width=28)
        self.company_cb.grid(row=4, column=1, columnspan=3, sticky="ew", padx=(8, 0), pady=(8, 0))
        self.company_cb.bind('<<ComboboxSelected>>', self._on_company_selected)

        ttk.Label(opts, text="Multiplo").grid(row=5, column=0, sticky="w")
        self.multiplier_sp = ttk.Spinbox(opts, from_=1, to=100000, width=8)
        self.multiplier_sp.set("80")
        self.multiplier_sp.grid(row=5, column=1, sticky="w", padx=(8, 0))
        try:
            self.default_multiplier = 80
        except Exception:
            self.default_multiplier = 80
        self.multiplier_sp.configure(command=self._on_multiplier_change)
        self.multiplier_sp.bind('<FocusOut>', lambda e: self._on_multiplier_change())
        self.multiplier_sp.bind('<Return>', lambda e: self._on_multiplier_change())

        btns = ttk.Frame(opts)
        btns.grid(row=6, column=0, columnspan=6, sticky="ew", pady=(12, 0))
        btns.columnconfigure(0, weight=1)
        ttk.Button(btns, text="Analizar", style="Accent.TButton", command=self.analyze_all).grid(row=0, column=0, sticky="ew")
        self.export_excel_btn = ttk.Button(btns, text="Exportar Excel", command=self.export_excel, state="disabled")
        self.export_excel_btn.grid(row=1, column=0, sticky="ew", pady=(8, 0))
        self.export_csv_btn = ttk.Button(btns, text="Exportar CSV", command=self.export_csv, state="disabled")
        self.export_csv_btn.grid(row=1, column=1, sticky="ew", padx=(8, 0), pady=(8, 0))
        self.clear_btn = ttk.Button(btns, text="Limpiar", command=self.clear_results)
        self.clear_btn.grid(row=2, column=0, sticky="ew", pady=(8, 0))
        self.report_btn = ttk.Button(btns, text="Generar reporte mensual", command=self.generate_report, state="disabled")
        self.report_btn.grid(row=2, column=1, sticky="ew", padx=(8, 0), pady=(8, 0))

        right = ttk.Labelframe(body, text="Registro y resultados", style="Section.TLabelframe")
        right.grid(row=0, column=1, sticky="nsew")
        right.rowconfigure(0, weight=1)
        right.columnconfigure(0, weight=1)

        self.info_text = tk.Text(right, height=20, wrap="word", font=("Consolas", 10))
        yscroll = ttk.Scrollbar(right, orient="vertical", command=self.info_text.yview)
        self.info_text.configure(yscrollcommand=yscroll.set)
        self.info_text.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")

        body.rowconfigure(0, weight=1)
        body.columnconfigure(1, weight=1)

    def _build_statusbar(self):
        bar = ttk.Frame(self.root, padding=(12, 6))
        bar.grid(row=2, column=0, sticky="ew")
        bar.columnconfigure(0, weight=1)
        self.status_label = ttk.Label(bar, text="Listo")
        self.status_label.grid(row=0, column=0, sticky="w")
        self.progress = ttk.Progressbar(bar, mode="indeterminate", length=160)
        self.progress.grid(row=0, column=1, sticky="e")

    def set_busy(self, busy: bool, msg: str = ""):
        if busy:
            self.status_label.config(text=msg or "Procesando…")
            try:
                self.progress.start(12)
            except Exception:
                pass
        else:
            try:
                self.progress.stop()
            except Exception:
                pass
            self.status_label.config(text=msg or "Listo")

    def append_info(self, text):
        self.info_text.configure(state="normal")
        self.info_text.insert("end", text + "\n")
        self.info_text.see("end")
        self.info_text.configure(state="disabled")

    def show_error(self, message: str):
        try:
            messagebox.showerror("Error", message)
        finally:
            try:
                self.append_info(f"[ERROR] {message}")
            except Exception:
                pass

    def browse_folder(self):
        folder = filedialog.askdirectory(title="Selecciona carpeta con archivos")
        if folder:
            self.folder_path.delete(0, tk.END)
            self.folder_path.insert(0, folder)

    def _sanitize_time_inputs(self):
        def to_int(v, default):
            try:
                return int(str(v).strip())
            except Exception:
                return default
        def clamp(v, lo, hi):
            return max(lo, min(hi, v))
        sh = clamp(to_int(self.start_hour.get(), 0), 0, 23)
        sm = clamp(to_int(self.start_min.get(), 0), 0, 59)
        eh = clamp(to_int(self.end_hour.get(), 23), 0, 23)
        em = clamp(to_int(self.end_min.get(), 59), 0, 59)
        try:
            self.start_hour.set(f"{sh:02d}")
            self.start_min.set(f"{sm:02d}")
            self.end_hour.set(f"{eh:02d}")
            self.end_min.set(f"{em:02d}")
        except Exception:
            pass
        return sh, sm, eh, em

    def analyze_folder(self):
        self._run_analysis("csv")

    def analyze_all(self):
        self._run_analysis("csv")

    def _run_analysis(self, file_type="csv"):
        folder_path = self.folder_path.get()
        if not folder_path or not os.path.exists(folder_path):
            messagebox.showerror("Error", "Selecciona una carpeta válida")
            return

        sdate = self.start_date.get_date()
        edate = self.end_date.get_date()
        sh, sm, eh, em = self._sanitize_time_inputs()
        start_dt = datetime(sdate.year, sdate.month, sdate.day, sh, sm)
        end_dt = datetime(edate.year, edate.month, edate.day, eh, em)
        if end_dt < start_dt:
            messagebox.showerror("Error", "El fin debe ser posterior al inicio")
            return

        resolution = self.resolution.get()

        self.info_text.configure(state="normal")
        self.info_text.delete("1.0", "end")
        self.info_text.configure(state="disabled")
        self.set_busy(True, "Procesando…")

        def progress_cb(msg: str):
            self.root.after(0, lambda: self.append_info(msg))

        def month_span(s, e):
            y, m = s.year, s.month
            while (y < e.year) or (y == e.year and m <= e.month):
                yield y, m
                m += 1
                if m > 12:
                    m = 1
                    y += 1

        def worker():
            try:
                months = list(month_span(sdate, edate))
                monthly_dfs, all_details = [], []
                total_files = total_processed = total_errors = 0
                last_folder = folder_path

                for (yy, mm) in months:
                    if (yy, mm) == (sdate.year, sdate.month) and (yy, mm) == (edate.year, edate.month):
                        s_t = f"{sh:02d}:{sm:02d}"
                        e_t = f"{eh:02d}:{em:02d}"
                    elif (yy, mm) == (sdate.year, sdate.month):
                        s_t, e_t = f"{sh:02d}:{sm:02d}", "23:59"
                    elif (yy, mm) == (edate.year, edate.month):
                        s_t, e_t = "00:00", f"{eh:02d}:{em:02d}"
                    else:
                        s_t, e_t = "00:00", "23:59"

                    ok, msg, results = self.csv_processor.analyze_folder(
                        Path(folder_path), mes_usuario=mm, año_usuario=yy,
                        start_time=s_t, end_time=e_t, progress_cb=progress_cb
                    )
                    if ok and getattr(self.csv_processor, "combined_df", None) is not None:
                        monthly_dfs.append(self.csv_processor.combined_df.copy())
                    if results:
                        all_details.extend(results.get("file_details", []))
                        total_files += results.get("total_files", 0)
                        total_processed += results.get("processed_files", 0)
                        total_errors += results.get("error_files", 0)
                        last_folder = results.get("folder", last_folder)

                if not monthly_dfs:
                    self.root.after(0, lambda: self.set_busy(False, "Sin datos"))
                    self.root.after(0, lambda: self.append_info("No se generaron datos"))
                    return

                combined = pd.concat(monthly_dfs, ignore_index=True)

                if "timestamp" in combined.columns and not pd.api.types.is_datetime64_any_dtype(combined["timestamp"]):
                    combined["timestamp"] = pd.to_datetime(
                        combined["timestamp"].astype(str).str.strip(),
                        errors="coerce", dayfirst=True
                    )
                if "timestamp" in combined.columns:
                    before_rows = combined.shape[0]
                    combined = combined[pd.notna(combined["timestamp"])].copy()
                    after_parse_rows = combined.shape[0]
                else:
                    before_rows = combined.shape[0]
                    after_parse_rows = before_rows

                if "timestamp" in combined.columns:
                    combined = combined[(combined["timestamp"] >= start_dt) & (combined["timestamp"] <= end_dt)]
                after_filter_rows = combined.shape[0]

                try:
                    kwh_nonnull = int(pd.notna(combined["kwh"]).sum()) if "kwh" in combined.columns else 0
                    kvar_nonnull = int(pd.notna(combined["kvarh"]).sum()) if "kvarh" in combined.columns else 0
                    kwh_sum_preview = float(pd.to_numeric(combined.get("kwh"), errors="coerce").sum()) if "kwh" in combined.columns else 0.0
                    kvar_sum_preview = float(pd.to_numeric(combined.get("kvarh"), errors="coerce").sum()) if "kvarh" in combined.columns else 0.0
                    self.root.after(0, lambda: self.append_info(f"Preview energía → kwh_nonnull={kwh_nonnull} kvar_nonnull={kvar_nonnull} kwh_sum={kwh_sum_preview:.3f} kvar_sum={kvar_sum_preview:.3f}"))
                except Exception:
                    pass

                if resolution == "1h" and not combined.empty and "timestamp" in combined.columns:
                    combined["hour_ts"] = combined["timestamp"].dt.floor("h")
                    agg_cols = {c: "sum" for c in ["kwh", "kvarh"] if c in combined.columns}
                    grouped = combined.groupby(["company", "hour_ts"], as_index=False).agg(agg_cols)
                    combined = grouped.rename(columns={"hour_ts": "timestamp"})
                    combined = combined.sort_values(["company", "timestamp"]).reset_index(drop=True)

                self.csv_processor.combined_df = combined

                fmt = "%d/%m/%y %H:%M"
                agg = defaultdict(lambda: {
                    "filename": None, "rows": 0, "success": False,
                    "kwh_values": 0, "kvar_values": 0,
                    "start_date": None, "end_date": None, "error": None
                })
                for d in all_details:
                    name = d.get("filename")
                    if not name:
                        continue
                    a = agg[name]
                    a["filename"] = name
                    a["rows"] += int(d.get("rows", 0))
                    a["success"] = a["success"] or bool(d.get("success", False))
                    a["kwh_values"] += int(d.get("kwh_values", 0))
                    a["kvar_values"] += int(d.get("kvar_values", 0))
                dedup_details = list(agg.values())
                for a in dedup_details:
                    a["start_date"] = start_dt.strftime(fmt)
                    a["end_date"] = end_dt.strftime(fmt)

                total_files_u = len(dedup_details)
                processed_u = sum(1 for x in dedup_details if x.get("success"))
                error_u = total_files_u - processed_u

                results_agg = {
                    "folder": last_folder,
                    "total_files": total_files_u,
                    "processed_files": processed_u,
                    "error_files": error_u,
                    "date_range": {
                        "start": start_dt.strftime("%d/%m/%y %H:%M"),
                        "end": end_dt.strftime("%d/%m/%y %H:%M")
                    },
                    "combined_stats": {
                        "total_rows": int(combined.shape[0]),
                        "total_columns": int(combined.shape[1]),
                        "total_kwh_values": int(pd.notna(combined["kwh"]).sum()) if "kwh" in combined.columns else 0,
                        "total_kvar_values": int(pd.notna(combined["kvarh"]).sum()) if "kvarh" in combined.columns else 0,
                        "sum_kwh": float(pd.to_numeric(combined.get("kwh"), errors="coerce").sum()) if "kwh" in combined.columns else 0.0,
                        "sum_kvarh": float(pd.to_numeric(combined.get("kvarh"), errors="coerce").sum()) if "kvarh" in combined.columns else 0.0,
                        "resolution": resolution,
                        "rows_before_parse": before_rows,
                        "rows_after_parse": after_parse_rows,
                        "rows_after_filter": after_filter_rows
                    },
                    "file_details": dedup_details,
                    "errors": []
                }
                self.root.after(0, lambda: self.on_analysis_done(True, "Procesamiento CSV completado", results_agg))
            except Exception as e:
                self.root.after(0, lambda: self.show_error(str(e)))
            finally:
                self.root.after(0, lambda: self.set_busy(False, "Listo"))

        threading.Thread(target=worker, daemon=True).start()

    def clear_results(self):
        self.info_text.configure(state="normal")
        self.info_text.delete("1.0", "end")
        self.info_text.configure(state="disabled")
        self.export_excel_btn.configure(state="disabled")
        self.export_csv_btn.configure(state="disabled")
        self.append_info("Panel limpiado.")
        self.last_results = None
        self.csv_processor.combined_df = None

    def populate_companies(self):
        df = getattr(self.csv_processor, "combined_df", None)
        if df is None or df.empty:
            self.company_cb.configure(state="disabled", values=[])
            self.report_btn.configure(state="disabled")
            return
        if "company" in df.columns:
            companies = sorted([str(x) for x in df["company"].dropna().unique().tolist()])
        else:
            companies = ["General"]
            df["company"] = "General"
        self.company_cb.configure(state="readonly", values=companies)
        if not self.company_cb.get():
            self.company_cb.set(companies[0])
        self._on_company_selected()
        self.report_btn.configure(state="normal")

    def _on_company_selected(self, event=None):
        company = self.company_cb.get()
        if not company:
            return
        val = self.company_multipliers.get(company)
        if val is None:
            val = getattr(self, 'default_multiplier', 80)
        try:
            self.multiplier_sp.set(str(int(val)))
        except Exception:
            self.multiplier_sp.set(str(val))

    def _on_multiplier_change(self):
        company = self.company_cb.get()
        if not company:
            return
        try:
            m = float(self.multiplier_sp.get())
        except Exception:
            return
        self.company_multipliers[company] = m

    def _hourly_aggregate(self, df):
        if "timestamp" not in df.columns:
            return df.iloc[0:0].copy()
        if not pd.api.types.is_datetime64_any_dtype(df["timestamp"]):
            df = df.copy()
            df["timestamp"] = pd.to_datetime(df["timestamp"].astype(str), errors="coerce", dayfirst=True)
            df = df[pd.notna(df["timestamp"])]
        df["hour_ts"] = df["timestamp"].dt.floor("h")
        agg_cols = {}
        if "kwh" in df.columns:
            agg_cols["kwh"] = "sum"
        if "kvarh" in df.columns:
            agg_cols["kvarh"] = "sum"
        if not agg_cols:
            return df.iloc[0:0].copy()
        g = df.groupby("hour_ts", as_index=True).agg(agg_cols)
        return g

    def compute_report_table(self, company: str, start_dt: datetime, end_dt: datetime, multiplo: float):
        df = getattr(self.csv_processor, "combined_df", None)
        if df is None or df.empty:
            return pd.DataFrame(), {"kwh": 0.0, "kvarh": 0.0}, {}
        if "company" in df.columns:
            df = df[df["company"].astype(str) == str(company)].copy()
        if "timestamp" in df.columns:
            if not pd.api.types.is_datetime64_any_dtype(df["timestamp"]):
                df["timestamp"] = pd.to_datetime(df["timestamp"].astype(str), errors="coerce", dayfirst=True)
                df = df[pd.notna(df["timestamp"])]
            df = df[(df["timestamp"] >= start_dt) & (df["timestamp"] <= end_dt)]
        hourly = self._hourly_aggregate(df)
        hourly_dict = hourly.to_dict("index")
        has_kwh = "kwh" in hourly.columns
        has_kvarh = "kvarh" in hourly.columns
        rows = []
        cur = datetime(start_dt.year, start_dt.month, start_dt.day)
        end_day = datetime(end_dt.year, end_dt.month, end_dt.day)
        while cur <= end_day:
            for h in range(24):
                ts = cur + timedelta(hours=h)
                kv = hourly_dict.get(ts)
                kwh = float(kv["kwh"]) if kv and has_kwh else 0.0
                kvarh = float(kv["kvarh"]) if kv and has_kvarh else 0.0
                kwh_scaled = kwh * multiplo
                kvarh_scaled = kvarh * multiplo
                rows.append({
                    "Fecha": format_es_date(cur.date()),
                    "Hora": h + 1,
                    "Kwh": round(kwh_scaled, 3), # Usar 3 decimales
                    "Kvarh": round(kvarh_scaled, 3), # Usar 3 decimales
                })
            cur += timedelta(days=1)
        report = pd.DataFrame(rows, columns=["Fecha", "Hora", "Kwh", "Kvarh"])
        totals = {
            "kwh": float(report["Kwh"].sum()),
            "kvarh": float(report["Kvarh"].sum()),
        }
        meta = {"company": company, "multiplo": multiplo}
        return report, totals, meta

    def generate_report(self):
        df = getattr(self.csv_processor, "combined_df", None)
        if df is None or df.empty:
            messagebox.showinfo("Reporte", "No hay datos para generar reporte.")
            return
        company = self.company_cb.get() or ("General" if "company" not in df.columns else str(df["company"].iloc[0]))
        try:
            m = float(self.multiplier_sp.get())
        except Exception:
            m = 80.0
        self.company_multipliers[company] = m
        sdate = self.start_date.get_date(); edate = self.end_date.get_date()
        sh, sm, eh, em = self._sanitize_time_inputs()
        start_dt = datetime(sdate.year, sdate.month, sdate.day, 0, 0)
        end_dt = datetime(edate.year, edate.month, edate.day, 23, 59)
        report_df, totals, meta = self.compute_report_table(company, start_dt, end_dt, m)
        if report_df.empty:
            messagebox.showinfo("Reporte", "No se generaron filas para el rango seleccionado.")
            return
        self.last_report = {"df": report_df, "totals": totals, "meta": meta}
        self.show_report_window(report_df, totals, meta)

    def show_report_window(self, report_df: pd.DataFrame, totals: dict, meta: dict):
        win = tk.Toplevel(self.root)
        win.title(f"Reporte mensual - {meta.get('company','')}")
        win.geometry("700x680")
        top = ttk.Frame(win, padding=12)
        top.pack(side="top", fill="x")
        ttk.Label(top, text=f"Multiplo → {int(meta.get('multiplo', 0))}", font=("Segoe UI", 12, "bold")).pack(side="left")
        sep = ttk.Frame(top); sep.pack(side="left", expand=True, fill="x")
        ttk.Label(top, text=f"{totals['kwh']:,.3f}", font=("Segoe UI", 14, "bold"), foreground="#1b4f72").pack(side="left", padx=(8, 16))
        ttk.Label(top, text=f"{totals['kvarh']:,.3f}", font=("Segoe UI", 14, "bold"), foreground="#1b4f72").pack(side="left")

        cols = ("Fecha", "Hora", "Kwh", "Kvarh")
        tree = ttk.Treeview(win, columns=cols, show="headings", height=24)
        for c, w in [("Fecha", 120), ("Hora", 60), ("Kwh", 120), ("Kvarh", 120)]:
            tree.heading(c, text=c)
            tree.column(c, width=w, anchor="e" if c in ("Hora", "Kwh", "Kvarh") else "w")
        for _, r in report_df.iterrows():
            tree.insert("", "end", values=(r["Fecha"], f"{int(r['Hora'])}", f"{r['Kwh']:.3f}", f"{r['Kvarh']:.3f}"))
        tree.pack(side="top", fill="both", expand=True, padx=12, pady=8)

        btnf = ttk.Frame(win, padding=12)
        btnf.pack(side="bottom", fill="x")
        ttk.Button(btnf, text="Exportar reporte a Excel", command=self.export_report_excel).pack(side="right")

    def export_report_excel(self):
        if not self.last_report:
            messagebox.showinfo("Exportar", "No hay reporte para exportar.")
            return
        report_df = self.last_report["df"]
        totals = self.last_report["totals"]
        meta = self.last_report["meta"]
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            title="Guardar reporte mensual"
        )
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte"

            ws["A1"] = "Multiplo →"
            ws["B1"] = int(meta.get("multiplo", 0))
            ws["C1"] = "Kwh"
            ws["D1"] = "Kvarh"
            ws["C2"] = totals["kwh"]
            ws["D2"] = totals["kvarh"]
            ws["C2"].number_format = "#,##0.000"
            ws["D2"].number_format = "#,##0.000"
            ws["A1"].font = Font(bold=True, size=12)
            ws["C2"].font = Font(bold=True, size=14)
            ws["D2"].font = Font(bold=True, size=14)

            headers = ["Fecha", "Hora", "Kwh", "Kvarh"]
            ws.append([])
            ws.append(headers)
            hdr_fill = PatternFill("solid", fgColor="D9EAF7")
            for col, h in enumerate(headers, start=1):
                cell = ws.cell(row=4, column=col)
                cell.fill = hdr_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

            thin = Side(style="thin", color="999999")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)
            start_row = 5
            for idx, r in report_df.iterrows():
                row = start_row + idx
                ws.cell(row=row, column=1, value=r["Fecha"])
                ws.cell(row=row, column=2, value=int(r["Hora"]))
                c3 = ws.cell(row=row, column=3, value=float(r["Kwh"]))
                c4 = ws.cell(row=row, column=4, value=float(r["Kvarh"]))
                c3.number_format = "#,##0.000"
                c4.number_format = "#,##0.000"
                for col in range(1, 5):
                    ws.cell(row=row, column=col).border = border
                ws.cell(row=row, column=3).fill = PatternFill("solid", fgColor="E9F5FE")
                ws.cell(row=row, column=4).fill = PatternFill("solid", fgColor="E9F5FE")

            ws.column_dimensions["A"].width = 14
            ws.column_dimensions["B"].width = 8
            ws.column_dimensions["C"].width = 14
            ws.column_dimensions["D"].width = 14

            wb.save(path)
            messagebox.showinfo("Exportar", f"Archivo guardado:\n{path}")
        except Exception as e:
            self.show_error(str(e))

    def export_excel(self):
        df = getattr(self.csv_processor, "combined_df", None)
        if df is None or df.empty:
            messagebox.showinfo("Exportar", "No hay datos para exportar.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            title="Guardar Excel (multi-hoja)"
        )
        if not path:
            return

        selected_company = self.company_cb.get() or None
        try:
            selected_multiplo = float(self.multiplier_sp.get())
        except Exception:
            selected_multiplo = None

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

            data = df.copy()
            if "timestamp" in data.columns and not data["timestamp"].isna().all():
                if pd.api.types.is_datetime64_any_dtype(data["timestamp"]):
                    data["timestamp"] = data["timestamp"].dt.strftime("%d/%m/%Y %H:%M:%S")
                else:
                    ts = pd.to_datetime(data["timestamp"].astype(str), errors="coerce", dayfirst=True)
                    data["timestamp"] = ts.dt.strftime("%d/%m/%Y %H:%M:%S")

            wb = Workbook()
            
            # Estilos sin fondo amarillo estridente (para coincidir con imagen 2 que es minimalista)
            cyan_fill = PatternFill("solid", fgColor="D9EAF7") # Cabecera suave
            light_fill = PatternFill("solid", fgColor="E9F5FE") 
            thin = Side(style="thin", color="000000")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            ws_total = wb.active
            ws_total.title = "total"

            total_headers = ["No.", "Cliente", "Multiplo", "KWh", "KVARh", "KW"]
            ws_total.append(total_headers)
            
            for c in range(1, len(total_headers) + 1):
                cell = ws_total.cell(row=1, column=c)
                cell.fill = cyan_fill 
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")
                cell.border = border

            companies = sorted([str(x) for x in data["company"].dropna().unique()]) if "company" in data.columns else ["General"]
            if "company" not in data.columns:
                data["company"] = "General"

            used_titles = set([ws_total.title])
            def unique_title(base: str) -> str:
                t = str(base)[:31]
                if t not in used_titles:
                    used_titles.add(t); return t
                i = 2
                while True:
                    cand = (str(base)[:31-len(str(i))-1] + f" {i}") if len(str(base)) >= 31 else f"{str(base)} {i}"
                    cand = cand[:31]
                    if cand not in used_titles:
                        used_titles.add(cand); return cand
                    i += 1

            sheet_by_company = {}
            total_rows = []
            
            # CORRECCIÓN: Se restaura el cálculo estándar de demanda (factor de multiplicación 4 para resoluciones de 15 minutos)
            factor_demanda = 1 if self.resolution.get() == "1h" else 4

            for idx, company in enumerate(companies, start=1):
                m = float(self.company_multipliers.get(company, getattr(self, 'default_multiplier', 80)))
                if selected_company and company == selected_company and (selected_multiplo is not None):
                    m = float(selected_multiplo)
                    self.company_multipliers[company] = m
                cdf = data[data["company"].astype(str) == company].copy()

                kwh_sum = float(pd.to_numeric(cdf.get("kwh"), errors="coerce").sum()) if "kwh" in cdf.columns else 0.0
                kvar_sum = float(pd.to_numeric(cdf.get("kvarh"), errors="coerce").sum()) if "kvarh" in cdf.columns else 0.0
                kwh_total = kwh_sum * m
                kvar_total = kvar_sum * m

                sheet_name = unique_title(company)
                ws = wb.create_sheet(title=sheet_name)
                sheet_by_company[company] = sheet_name

                cols = []
                for name in ["timestamp", "company", "kwh", "kvarh"]:
                    if name in cdf.columns and name not in cols:
                        cols.append(name)
                for name in cdf.columns:
                    if name not in cols:
                        cols.append(name)

                from openpyxl.utils import get_column_letter
                kwh_col_idx = cols.index("kwh") + 1 if "kwh" in cols else None
                kvar_col_idx = cols.index("kvarh") + 1 if "kvarh" in cols else None
                kwh_letter = get_column_letter(kwh_col_idx) if kwh_col_idx else "C"
                kvar_letter = get_column_letter(kvar_col_idx) if kvar_col_idx else "D"

                ws["A1"] = "Multiplo →"; ws["A1"].font = Font(bold=True, size=12)
                ws["B1"] = int(m)
                
                kwh_hdr_cell = ws.cell(row=1, column=kwh_col_idx or 3)
                kwh_hdr_cell.value = kwh_total; kwh_hdr_cell.number_format = "#,##0.000"; kwh_hdr_cell.font = Font(bold=True, size=14)
                ws.cell(row=2, column=kwh_col_idx or 3, value="Kwh").font = Font(bold=True, size=10)
                
                kvar_hdr_cell = ws.cell(row=1, column=kvar_col_idx or 4)
                kvar_hdr_cell.value = kvar_total; kvar_hdr_cell.number_format = "#,##0.000"; kvar_hdr_cell.font = Font(bold=True, size=14)
                ws.cell(row=2, column=kvar_col_idx or 4, value="Kvarh").font = Font(bold=True, size=10)

                ws.append([])
                ws.append(cols)
                for col_idx_h, h in enumerate(cols, start=1):
                    cell = ws.cell(row=4, column=col_idx_h)
                    cell.fill = cyan_fill
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center")

                start_row = 5
                for ridx, (_, row) in enumerate(cdf.iterrows(), start=start_row):
                    for cidx, col_name in enumerate(cols, start=1):
                        val = row.get(col_name)
                        cell = ws.cell(row=ridx, column=cidx, value=val if pd.notna(val) else None)
                        cell.border = border
                        if col_name.lower() in ("kwh", "kvarh"):
                            cell.number_format = "#,##0.000"
                            cell.fill = light_fill

                ws.column_dimensions["A"].width = max(14, min(28, len(str(cols[0])) + 6)) if cols else 14
                ws.column_dimensions["B"].width = 10
                ws.column_dimensions["C"].width = 16
                ws.column_dimensions["D"].width = 16

                last_row = start_row + max(len(cdf), 1) - 1
                esc = sheet_name.replace("'", "''")
                
                # CORRECCIÓN: La fórmula de Excel vuelve a ser =MAX(Rango) * Factor_Demanda * Múltiplo
                if len(cdf) <= 0:
                    kw_formula = 0
                else:
                    kw_formula = f"=MAX('{esc}'!${kwh_letter}${start_row}:${kwh_letter}${last_row})*{factor_demanda}*{m}"
                
                total_rows.append((idx, company, m, kwh_total, kvar_total, kw_formula, esc,
                                   start_row, last_row, kwh_letter, kvar_letter))

            for r_idx, (no, company, m, kwh_t, kvar_t, kw_formula, esc, srow, lrow, kl, vl) in enumerate(total_rows, start=2):
                c1 = ws_total.cell(row=r_idx, column=1, value=no)
                c_name = ws_total.cell(row=r_idx, column=2, value=company)
                sheet_name = sheet_by_company[company]
                esc = sheet_name.replace("'", "''")
                c_name.hyperlink = f"#'{esc}'!A1"
                c_name.font = Font(color="0563C1", underline="single")
                c3 = ws_total.cell(row=r_idx, column=3, value=int(m))
                
                c4 = ws_total.cell(row=r_idx, column=4)
                c4.value = f"=IF(ISNUMBER('{esc}'!${kl}$1), '{esc}'!${kl}$1, SUM('{esc}'!${kl}${srow}:${kl}${lrow})*$C{r_idx})"
                c4.number_format = "#,##0.000"

                c5 = ws_total.cell(row=r_idx, column=5)
                c5.value = f"=IF(ISNUMBER('{esc}'!${vl}$1), '{esc}'!${vl}$1, SUM('{esc}'!${vl}${srow}:${vl}${lrow})*$C{r_idx})"
                c5.number_format = "#,##0.000"
                
                c6 = ws_total.cell(row=r_idx, column=6)
                if isinstance(kw_formula, str):
                    c6.value = kw_formula
                else:
                    c6.value = 0
                c6.number_format = "#,##0.000" # Restaurado a 3 decimales
                
                for col_idx in range(1, 7):
                    ws_total.cell(row=r_idx, column=col_idx).border = border
                    ws_total.cell(row=r_idx, column=col_idx).alignment = Alignment(horizontal="center" if col_idx in (1, 3) else "right" if col_idx > 3 else "left")

            ws_total.column_dimensions["A"].width = 6
            ws_total.column_dimensions["B"].width = 34
            ws_total.column_dimensions["C"].width = 10
            ws_total.column_dimensions["D"].width = 16
            ws_total.column_dimensions["E"].width = 16
            ws_total.column_dimensions["F"].width = 16

            wb.save(path)
            messagebox.showinfo("Exportar", f"Excel exportado: {path}")
        except Exception as e:
            self.show_error(str(e))

    def export_csv(self):
        df = getattr(self.csv_processor, "combined_df", None)
        if df is None or df.empty:
            messagebox.showinfo("Exportar", "No hay datos para exportar.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv")],
            title="Guardar CSV combinado"
        )
        if not path:
            return
        ok, msg = self.csv_processor.export_combined_csv(path)
        if ok:
            messagebox.showinfo("Exportar", msg)
        else:
            messagebox.showerror("Exportar", msg)

    def on_analysis_done(self, ok: bool, msg: str, results: dict):
        self.append_info(msg)
        if ok and getattr(self.csv_processor, "combined_df", None) is not None:
            self.last_results = results
            if hasattr(self, "export_excel_btn"):
                self.export_excel_btn.configure(state="normal")
            if hasattr(self, "export_csv_btn"):
                self.export_csv_btn.configure(state="normal")
            self.populate_companies()
            cs = results.get("combined_stats", {})
            self.append_info(f"Filas: {cs.get('total_rows', 0)}  Columnas: {cs.get('total_columns', 0)}  Resolución: {cs.get('resolution', '')}")
        else:
            self.append_info("Sin resultados para exportar.")
            self.company_cb.configure(state="disabled")
            self.report_btn.configure(state="disabled")